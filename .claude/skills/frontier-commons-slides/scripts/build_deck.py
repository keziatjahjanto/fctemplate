#!/usr/bin/env python3
"""Build a Frontier Commons branded PowerPoint deck from a JSON outline.

Usage:
    python3 build_deck.py deck.json [-o out.pptx]

deck.json:
    {
      "title": "Deck title (file metadata)",
      "footer": "Frontier Commons",          # optional footer text on content slides
      "slides": [ { "layout": "title", ... }, ... ]
    }

Layouts and their fields (* = required):
    title       title*, subtitle, date
    agenda      title, items* [str]
    section     title*, kicker
    content     title*, tag, body, bullets [str]
    two_column  title*, tag, left* {heading, body, bullets}, right* {heading, body, bullets}
    cards       title*, tag, cards* [{title*, subtitle, text}]  (2-4)
    stats       title*, tag, stats* [{value*, label*, detail}]   (2-4)
    process     title*, tag, steps* [{title*, text}]             (3-6)
    quote       quote*, author, role
    image       title*, image*, tag, body, bullets, caption, side ("left"|"right")
    statement   text*, tag                                       (one big sentence)
    closing     title, subtitle, contact

Every slide accepts "notes" (speaker notes).
Text supports **highlight** (bold + brand accent colour) and *italic*.
Image paths are relative to the JSON file.
"""
import argparse
import json
import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "assets"

# ---- Brand tokens (Branding Guide, Jan 2026) ---------------------------------
ONYX = RGBColor(0x13, 0x14, 0x16)
DEEP_BLUE = RGBColor(0x0D, 0x14, 0x5F)
MIST = RGBColor(0xF3, 0xF4, 0xFA)
PERIWINKLE = RGBColor(0x9C, 0x90, 0xE6)
PAPAYA = RGBColor(0xFF, 0xEF, 0xD5)
AMBER = RGBColor(0xFA, 0xBC, 0x20)
BLUE_50 = RGBColor(0x1F, 0x2F, 0xE0)     # tonal step — text highlights on light backgrounds
PERI_82 = RGBColor(0xBC, 0xB4, 0xEE)     # tonal step — text highlights on dark backgrounds
ONYX_40 = RGBColor(0x5F, 0x64, 0x6D)     # tonal step — secondary text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

DISPLAY = "Inter"      # headers, headings
BODY = "DM Sans"       # section headers, body, captions

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.8)        # outer margin
CONTENT_W = W - 2 * M


# ---- Low-level helpers -------------------------------------------------------
def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color, rounded=False, radius=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    if rounded:
        shp.adjustments[0] = radius
    return shp


def oval(slide, x, y, d, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def tracking(run, thousandths):
    """Letter spacing in Figma-style units (1/1000 em), e.g. -20."""
    size = run.font.size.pt if run.font.size else 18
    run.font._element.set("spc", str(int(round(size * thousandths / 10))))


def add_runs(par, text, font, size, color, bold=False, italic=False, hl=None, track=0):
    """Add text to a paragraph, honouring **highlight** and *italic* markup."""
    tokens = re.split(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?\*)", str(text))
    for tok in tokens:
        if not tok:
            continue
        r = par.add_run()
        is_hl = tok.startswith("**") and tok.endswith("**")
        is_it = not is_hl and tok.startswith("*") and tok.endswith("*") and len(tok) > 1
        r.text = tok[2:-2] if is_hl else tok[1:-1] if is_it else tok
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = True if is_hl else bold
        f.italic = True if is_it else italic
        f.color.rgb = (hl or color) if is_hl else color
        if track:
            tracking(r, track)


def textbox(slide, x, y, w, h, text="", font=BODY, size=16, color=ONYX, bold=False, italic=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.15, hl=None, track=0, autofit=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else str(text).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        add_runs(p, ln, font, size, color, bold, italic, hl, track)
    if autofit:
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        for child in list(bodyPr):
            bodyPr.remove(child)
        bodyPr.append(bodyPr.makeelement(qn("a:normAutofit"), {}))
    return tb


def bullets(slide, x, y, w, h, items, size=18, color=ONYX, hl=BLUE_50, marker=DEEP_BLUE, gap=10):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        indent = int(Pt(size) * 1.1)
        pPr.set("marL", str(indent))
        pPr.set("indent", str(-indent))
        bu_clr = pPr.makeelement(qn("a:buClr"), {})
        srgb = bu_clr.makeelement(qn("a:srgbClr"), {"val": str(marker)})
        bu_clr.append(srgb)
        pPr.append(bu_clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        add_runs(p, item, BODY, size, color, hl=hl, track=-10)
    return tb


def picture(slide, path, x, y, w, h, radius_in=0.25):
    """Insert an image cropped to fill the box, with brand rounded corners."""
    im = Image.open(path).convert("RGBA")
    box_ratio = w / h
    iw, ih = im.size
    if iw / ih > box_ratio:
        nw = int(ih * box_ratio)
        im = im.crop(((iw - nw) // 2, 0, (iw - nw) // 2 + nw, ih))
    else:
        nh = int(iw / box_ratio)
        im = im.crop((0, (ih - nh) // 2, iw, (ih - nh) // 2 + nh))
    px_per_in = im.size[0] / (w / 914400)
    r = int(radius_in * px_per_in)
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, im.size[0] - 1, im.size[1] - 1], r, fill=255)
    im.putalpha(mask)
    out = Path(TMP_DIR) / f"img_{abs(hash((str(path), w, h)))}.png"
    im.save(out)
    return slide.shapes.add_picture(str(out), x, y, w, h)


def logo(slide, kind, x, y, height):
    """kind: icon-deep-blue | icon-mist | wordmark-horizontal-mist | wordmark-stacked-deep-blue ..."""
    name = f"logo-{kind}.png" if kind.startswith("icon") else f"{kind}.png"
    return slide.shapes.add_picture(str(ASSETS / name), x, y, height=height)


def tag(slide, x, y, text):
    """Amber pill label (accent — use sparingly)."""
    size = 12
    w = Emu(int(Pt(size) * 0.62 * len(text) + Inches(0.4)))
    h = Inches(0.36)
    shp = rect(slide, x, y, w, h, AMBER, rounded=True, radius=0.5)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, text, BODY, size, ONYX, bold=True)
    return h


def footer(slide, n, text, dark=False):
    color = MIST if dark else ONYX_40
    logo(slide, "icon-mist" if dark else "icon-deep-blue", W - M - Inches(0.32), H - Inches(0.62), Inches(0.32))
    textbox(slide, M, H - Inches(0.58), Inches(6), Inches(0.3), f"{n}   {text}" if text else str(n),
            font=BODY, size=10, color=color)


def header(slide, s, top=Inches(0.7), color=DEEP_BLUE, size=36):
    """Optional amber tag + slide title. Returns y where content can start."""
    y = top
    if s.get("tag"):
        tag(slide, M, y, s["tag"])
        y += Inches(0.55)
    textbox(slide, M, y, CONTENT_W, Inches(1.2), s.get("title", ""), font=DISPLAY, size=size,
            color=color, bold=True, track=-10, line=1.05, hl=BLUE_50, anchor=MSO_ANCHOR.TOP)
    lines = max(1, len(s.get("title", "")) // 48 + 1)
    return y + Pt(size) * 1.15 * lines + Inches(0.45)


# ---- Layouts -----------------------------------------------------------------
def l_title(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, DEEP_BLUE)
    textbox(sl, M, Inches(1.3), Inches(10), Inches(2.6), s["title"], font=DISPLAY, size=66, color=MIST,
            bold=True, line=1.0, anchor=MSO_ANCHOR.BOTTOM, hl=PERI_82)
    if s.get("subtitle"):
        textbox(sl, M, Inches(4.15), Inches(10), Inches(0.8), s["subtitle"].upper(), font=DISPLAY, size=24,
                color=MIST, bold=True, track=-10)
    rect(sl, M, Inches(6.45), CONTENT_W, Pt(1), MIST)
    if s.get("date"):
        textbox(sl, M, Inches(6.65), Inches(5), Inches(0.4), s["date"], font=BODY, size=14, color=MIST)
    logo(sl, "wordmark-horizontal-mist", W - M - Inches(2.9), Inches(6.68), Inches(0.37))
    return sl


def l_section(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, PERIWINKLE)
    line_y = int(H * 0.786)
    rect(sl, 0, line_y, W, Pt(1.25), MIST)
    size = 110 if len(s["title"]) <= 14 else 80
    # Box bottom sits one descender below the line so the baseline rests on it (brand signature).
    desc = int(Pt(size) * 0.25)
    textbox(sl, M, line_y + desc - Inches(3), CONTENT_W, Inches(3), s["title"], font=DISPLAY, size=size,
            color=MIST, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM, line=0.95, track=-20)
    if s.get("kicker"):
        textbox(sl, M, Inches(0.7), Inches(8), Inches(0.5), s["kicker"], font=BODY, size=16, color=MIST)
    logo(sl, "icon-deep-blue", M, H - Inches(0.85), Inches(0.4))
    return sl


def l_content(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, MIST)
    y = header(sl, s)
    rect(sl, M, y - Inches(0.25), Inches(0.9), Pt(3), DEEP_BLUE)
    avail = H - y - Inches(1.0)
    if s.get("body"):
        tb = textbox(sl, M, y, Inches(9.5), Inches(1.4), s["body"], font=BODY, size=20, color=ONYX,
                     line=1.4, track=-10, hl=BLUE_50)
        y += Inches(0.5) + Pt(20) * 1.4 * (len(s["body"]) // 85 + 1)
        avail = H - y - Inches(1.0)
    if s.get("bullets"):
        bullets(sl, M, y, Inches(10.5), avail, s["bullets"], size=20)
    footer(sl, n, deck.get("footer"))
    return sl


def _column(sl, x, y, w, h, col):
    if col.get("heading"):
        textbox(sl, x, y, w, Inches(0.5), col["heading"], font=BODY, size=22, color=DEEP_BLUE, bold=True, track=-10)
        y += Inches(0.6)
    if col.get("body"):
        textbox(sl, x, y, w, Inches(1.5), col["body"], font=BODY, size=17, color=ONYX, line=1.4, track=-10, hl=BLUE_50)
        y += Pt(17) * 1.4 * (len(col["body"]) // 55 + 1) + Inches(0.25)
    if col.get("bullets"):
        bullets(sl, x, y, w, h, col["bullets"], size=17)


def l_two_column(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, MIST)
    y = header(sl, s)
    gap = Inches(0.6)
    cw = (CONTENT_W - gap) / 2
    h = H - y - Inches(1)
    rect(sl, M, y, Emu(int(cw)), h, WHITE, rounded=True, radius=0.06)
    rect(sl, M + Emu(int(cw + gap)), y, Emu(int(cw)), h, WHITE, rounded=True, radius=0.06)
    pad = Inches(0.4)
    _column(sl, M + pad, y + pad, Emu(int(cw - 2 * pad)), h - 2 * pad, s["left"])
    _column(sl, M + Emu(int(cw + gap)) + pad, y + pad, Emu(int(cw - 2 * pad)), h - 2 * pad, s["right"])
    footer(sl, n, deck.get("footer"))
    return sl


def l_cards(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, WHITE)
    y = header(sl, s)
    cards = s["cards"][:4]
    gap = Inches(0.35)
    cw = (CONTENT_W - gap * (len(cards) - 1)) / len(cards)
    longest = max(len(c.get("text", "")) for c in cards)
    chars_per_line = max(12, int(cw / Inches(1)) * 9)
    ch = min(H - y - Inches(1.0), Inches(2.6) + Pt(15) * 1.4 * (longest // chars_per_line + 2))
    y = y + (H - y - Inches(1.0) - ch) / 2  # centre the row vertically
    y = Emu(int(y))
    for i, c in enumerate(cards):
        x = M + Emu(int((cw + gap) * i))
        rect(sl, x, y, Emu(int(cw)), ch, PAPAYA, rounded=True, radius=0.1)
        pad = Inches(0.3)
        iw = Emu(int(cw - 2 * pad))
        textbox(sl, x + pad, y + Inches(0.45), iw, Inches(0.5), c["title"], font=DISPLAY, size=22, color=ONYX,
                bold=True, align=PP_ALIGN.CENTER)
        uw = Inches(1.9)
        rect(sl, x + Emu(int((cw - uw) / 2)), y + Inches(1.0), uw, Pt(2.5), AMBER)
        ty = y + Inches(1.2)
        if c.get("subtitle"):
            textbox(sl, x + pad, ty, iw, Inches(0.8), c["subtitle"], font=DISPLAY, size=18, color=ONYX,
                    align=PP_ALIGN.CENTER)
            ty += Inches(0.85)
        if c.get("text"):
            textbox(sl, x + pad, ty + Inches(0.15), iw, ch - (ty - y) - Inches(0.4), c["text"], font=BODY, size=15,
                    color=ONYX, align=PP_ALIGN.CENTER, line=1.4, hl=DEEP_BLUE)
    footer(sl, n, deck.get("footer"))
    return sl


def l_stats(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, MIST)
    y = header(sl, s)
    stats = s["stats"][:4]
    gap = Inches(0.5)
    cw = (CONTENT_W - gap * (len(stats) - 1)) / len(stats)
    top = y + Inches(0.4)
    for i, st in enumerate(stats):
        x = M + Emu(int((cw + gap) * i))
        iw = Emu(int(cw))
        textbox(sl, x, top, iw, Inches(1.4), st["value"], font=DISPLAY, size=72 if len(stats) < 4 else 60,
                color=DEEP_BLUE, bold=True, anchor=MSO_ANCHOR.BOTTOM, line=0.9, track=-30)
        rect(sl, x, top + Inches(1.6), Inches(0.9), Pt(4), AMBER)
        textbox(sl, x, top + Inches(1.85), iw, Inches(0.9), st["label"], font=DISPLAY, size=20, color=ONYX,
                line=1.2, track=-10)
        if st.get("detail"):
            textbox(sl, x, top + Inches(2.75), iw, Inches(1.2), st["detail"], font=BODY, size=14, color=ONYX_40,
                    line=1.4)
    footer(sl, n, deck.get("footer"))
    return sl


def l_process(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, PAPAYA)
    y = header(sl, s)
    steps = s["steps"][:6]
    k = len(steps)
    colw = CONTENT_W / k
    d = Inches(0.8)
    cy = y + Inches(0.8)
    rect(sl, M + Emu(int(colw / 2)), cy + d / 2 - Pt(1.5), Emu(int(colw * (k - 1))), Pt(3), DEEP_BLUE)
    for i, st in enumerate(steps):
        cx = M + Emu(int(colw * i + colw / 2))
        c = oval(sl, cx - d / 2, cy, d, PERIWINKLE)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, str(i + 1), DISPLAY, 22, ONYX, bold=True)
        tw = Emu(int(colw - Inches(0.25)))
        textbox(sl, cx - tw / 2, cy + d + Inches(0.3), tw, Inches(0.7), st["title"], font=BODY, size=17,
                color=DEEP_BLUE, bold=True, align=PP_ALIGN.CENTER, line=1.15)
        if st.get("text"):
            textbox(sl, cx - tw / 2, cy + d + Inches(1.0), tw, Inches(2), st["text"], font=BODY, size=14,
                    color=ONYX, align=PP_ALIGN.CENTER, line=1.4)
    footer(sl, n, deck.get("footer"))
    return sl


def l_agenda(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, MIST)
    s = dict(s, title=s.get("title", "Agenda"))
    y = header(sl, s)
    items = s["items"][:8]
    two = len(items) > 4
    per_col = (len(items) + 1) // 2 if two else len(items)
    cw = (CONTENT_W - Inches(0.6)) / 2 if two else CONTENT_W
    rowh = min(Inches(1.05), (H - y - Inches(1)) / per_col)
    d = Inches(0.6)
    for i, item in enumerate(items):
        col, row = divmod(i, per_col)
        x = M + Emu(int((cw + Inches(0.6)) * col))
        yy = y + Emu(int(rowh * row))
        c = oval(sl, x, yy, d, PERIWINKLE)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, str(i + 1), DISPLAY, 18, ONYX, bold=True)
        textbox(sl, x + d + Inches(0.3), yy, Emu(int(cw - d - Inches(0.3))), d, item, font=DISPLAY, size=22,
                color=ONYX, anchor=MSO_ANCHOR.MIDDLE, track=-20)
    footer(sl, n, deck.get("footer"))
    return sl


def l_quote(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, PAPAYA)
    textbox(sl, M, Inches(0.7), Inches(2), Inches(1.6), "“", font=DISPLAY, size=160, color=AMBER, bold=True,
            line=1.0)
    size = 34 if len(s["quote"]) < 160 else 26
    textbox(sl, M, Inches(2.1), Inches(10.5), Inches(3.2), s["quote"], font=BODY, size=size, color=DEEP_BLUE,
            italic=True, line=1.3, track=-20, anchor=MSO_ANCHOR.MIDDLE, hl=BLUE_50)
    rect(sl, M, Inches(5.55), Inches(1.2), Pt(4), AMBER)
    if s.get("author"):
        textbox(sl, M, Inches(5.8), Inches(8), Inches(0.4), s["author"], font=BODY, size=18, color=DEEP_BLUE, bold=True)
    if s.get("role"):
        textbox(sl, M, Inches(6.2), Inches(8), Inches(0.4), s["role"], font=BODY, size=14, color=ONYX_40)
    footer(sl, n, deck.get("footer"))
    return sl


def l_image(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, MIST)
    img_w = Inches(5.6)
    img_left = s.get("side", "left") == "left"
    ix = M if img_left else W - M - img_w
    tx = M + img_w + Inches(0.7) if img_left else M
    tw = CONTENT_W - img_w - Inches(0.7)
    picture(sl, DECK_DIR / s["image"], ix, Inches(0.7), img_w, H - Inches(1.6))
    y = Inches(1.3)
    if s.get("tag"):
        tag(sl, tx, y, s["tag"])
        y += Inches(0.6)
    textbox(sl, tx, y, tw, Inches(1.8), s["title"], font=DISPLAY, size=34, color=DEEP_BLUE, bold=True, line=1.05,
            track=-10, hl=BLUE_50)
    y += Pt(34) * 1.1 * (len(s["title"]) // 26 + 1) + Inches(0.35)
    if s.get("body"):
        textbox(sl, tx, y, tw, Inches(2), s["body"], font=BODY, size=17, color=ONYX, line=1.4, track=-10, hl=BLUE_50)
        y += Pt(17) * 1.4 * (len(s["body"]) // 48 + 1) + Inches(0.3)
    if s.get("bullets"):
        bullets(sl, tx, y, tw, H - y - Inches(1), s["bullets"], size=17)
    if s.get("caption"):
        textbox(sl, ix, H - Inches(0.8), img_w, Inches(0.3), s["caption"], font=BODY, size=11, color=ONYX_40)
    footer(sl, n, deck.get("footer"))
    return sl


def l_statement(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, DEEP_BLUE)
    y = Inches(2.0)
    if s.get("tag"):
        tag(sl, M, Inches(1.4), s["tag"])
    textbox(sl, M, y, Inches(11), Inches(3.8), s["text"], font=DISPLAY, size=48, color=MIST, bold=True, line=1.1,
            track=-20, hl=PERI_82, anchor=MSO_ANCHOR.TOP)
    footer(sl, n, deck.get("footer"), dark=True)
    return sl


def l_closing(prs, s, n, deck):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg(sl, DEEP_BLUE)
    lh = Inches(1.9)
    pic = logo(sl, "wordmark-stacked-mist", 0, Inches(1.2), lh)
    pic.left = int((W - pic.width) / 2)
    textbox(sl, M, Inches(3.65), CONTENT_W, Inches(1), s.get("title", "Thank you"), font=DISPLAY, size=48,
            color=MIST, bold=True, align=PP_ALIGN.CENTER)
    if s.get("subtitle"):
        textbox(sl, M, Inches(4.6), CONTENT_W, Inches(0.6), s["subtitle"], font=DISPLAY, size=22, color=PERI_82,
                align=PP_ALIGN.CENTER, track=-20)
    if s.get("contact"):
        textbox(sl, M, Inches(5.7), CONTENT_W, Inches(0.8), s["contact"], font=BODY, size=16, color=MIST,
                align=PP_ALIGN.CENTER, line=1.4)
    return sl


LAYOUTS = {
    "title": l_title, "section": l_section, "content": l_content, "two_column": l_two_column,
    "cards": l_cards, "stats": l_stats, "process": l_process, "agenda": l_agenda, "quote": l_quote,
    "image": l_image, "statement": l_statement, "closing": l_closing,
}

DECK_DIR = Path(".")
TMP_DIR = None


def main():
    global DECK_DIR, TMP_DIR
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()

    deck_path = Path(a.deck).resolve()
    DECK_DIR = deck_path.parent
    deck = json.loads(deck_path.read_text())
    out = Path(a.out) if a.out else deck_path.with_suffix(".pptx")

    import tempfile
    TMP_DIR = tempfile.mkdtemp()

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prs.core_properties.title = deck.get("title", "")
    prs.core_properties.author = deck.get("author", "Frontier Commons")

    for n, s in enumerate(deck["slides"], 1):
        layout = s.get("layout", "content")
        if layout not in LAYOUTS:
            sys.exit(f"Slide {n}: unknown layout '{layout}'. Options: {', '.join(LAYOUTS)}")
        try:
            sl = LAYOUTS[layout](prs, s, n, deck)
        except KeyError as e:
            sys.exit(f"Slide {n} ({layout}): missing required field {e}")
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]

    prs.save(out)
    print(f"✓ {out}  ({len(deck['slides'])} slides)")


if __name__ == "__main__":
    main()
