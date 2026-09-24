#!/usr/bin/env python3
"""Render every slide of a .pptx to PNG with macOS Quick Look, without opening PowerPoint.

Usage:
    python3 preview_deck.py deck.pptx [-o OUTDIR]

Writes OUTDIR/slide-NN.png for each slide and OUTDIR/sheet-N.png contact sheets (6 slides each),
then prints their paths. Quick Look hangs on slides with speaker notes, so the preview copies drop
the notes (the real deck is untouched). Quick Look also ignores chart bar order and number
formats; PowerPoint shows them correctly.
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation


def single_slide_copy(src, keep, dest):
    p = Presentation(src)
    ids = p.slides._sldIdLst
    for i, sid in reversed(list(enumerate(list(ids)))):
        if i != keep:
            p.part.drop_rel(sid.rId)
            ids.remove(sid)
    slide = p.slides[0]
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith("/notesSlide"):
            slide.part.drop_rel(rel.rId)
    p.save(dest)


def main():
    if sys.platform != "darwin":
        sys.exit("Quick Look previews need macOS. Elsewhere, open the deck and export a PDF instead.")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--outdir")
    a = ap.parse_args()
    src = Path(a.deck).resolve()
    out = Path(a.outdir or tempfile.mkdtemp(prefix="fc-preview-")).resolve()
    out.mkdir(parents=True, exist_ok=True)
    n = len(Presentation(str(src)).slides)
    pngs = []
    with tempfile.TemporaryDirectory() as tmp:
        for k in range(n):
            copy = Path(tmp) / f"slide-{k + 1:02d}.pptx"
            single_slide_copy(str(src), k, str(copy))
            try:
                subprocess.run(["qlmanage", "-t", "-s", "2400", "-o", tmp, str(copy)], capture_output=True, timeout=30)
            except subprocess.TimeoutExpired:
                subprocess.run(["pkill", "qlmanage"], capture_output=True)
            thumb = Path(tmp) / f"{copy.name}.png"
            if thumb.exists():
                dest = out / f"slide-{k + 1:02d}.png"
                shutil.move(str(thumb), dest)
                pngs.append(dest)
            else:
                print(f"! slide {k + 1} did not render in Quick Look")
    if not pngs:
        sys.exit("No slides rendered.")
    ims = [Image.open(p).convert("RGB") for p in pngs]
    w = 700
    h = int(ims[0].height * w / ims[0].width)
    for part in range(0, len(ims), 6):
        sel = ims[part:part + 6]
        sheet = Image.new("RGB", (2 * w + 10, ((len(sel) + 1) // 2) * (h + 10)), "white")
        for i, im in enumerate(sel):
            sheet.paste(im.resize((w, h)), ((i % 2) * (w + 10), (i // 2) * (h + 10)))
        sheet.save(out / f"sheet-{part // 6 + 1}.png")
        print(out / f"sheet-{part // 6 + 1}.png")
    print(f"{len(pngs)}/{n} slides in {out}")


if __name__ == "__main__":
    main()
